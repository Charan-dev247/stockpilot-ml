
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RAW_DATA_DIR = PROJECT_ROOT / "data" / "raw"


def main():
    workbooks = list(RAW_DATA_DIR.glob("*.xlsx"))

    if not workbooks:
        raise FileNotFoundError(
            f"No Excel workbook found in {RAW_DATA_DIR}"
        )

    workbook_path = workbooks[0]

    print("=" * 70)
    print("STOCKPILOT ML - DATASET INSPECTION")
    print("=" * 70)
    print(f"Workbook: {workbook_path.name}")
    print(f"File size: {workbook_path.stat().st_size / (1024 ** 2):.2f} MB")

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:
        print(f"\nSheets: {workbook.sheetnames}")

        for sheet in workbook.worksheets:
            print("\n" + "-" * 70)
            print(f"Sheet: {sheet.title}")
            print(f"Reported rows: {sheet.max_row}")
            print(f"Reported columns: {sheet.max_column}")

            header = next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                )
            )
            print(f"Header: {list(header)}")

            # Read only 1,000 rows for an initial inspection.
            sample = pd.read_excel(
                workbook_path,
                sheet_name=sheet.title,
                nrows=1000,
            )

            print("\nColumn data types (sample):")
            print(sample.dtypes.to_string())

            print("\nMissing values in the sample:")
            missing = sample.isna().sum()
            print(missing[missing > 0].to_string() or "None")

            print("\nFirst five records:")
            print(sample.head().to_string(index=False))

            if "InvoiceDate" in sample.columns:
                dates = pd.to_datetime(
                    sample["InvoiceDate"],
                    errors="coerce",
                )
                print("\nSample date range:")
                print(f"  Earliest: {dates.min()}")
                print(f"  Latest:   {dates.max()}")

    finally:
        workbook.close()

    print("\n" + "=" * 70)
    print("Initial inspection completed.")
    print("These are sample-level results, not full-dataset statistics.")
    print("=" * 70)


if __name__ == "__main__":
    main()
